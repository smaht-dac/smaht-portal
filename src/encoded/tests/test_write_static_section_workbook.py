import json
from pathlib import Path

import openpyxl
import pytest

from ..commands.write_static_section_workbook import (
    STATIC_SECTION_SHEET_NAME,
    STATIC_SECTION_WORKBOOK_FILENAME,
    StaticSectionWorkbookError,
    generate_workbook,
    get_output_path,
    get_static_section_columns,
    get_static_section_property_schemas,
    get_static_section_schema,
    load_config,
    validate_config,
    write_static_section_workbook,
)


@pytest.fixture
def static_section_schema():
    return {
        "required": ["identifier"],
        "properties": {
            "identifier": {
                "title": "Identifier",
                "type": "string",
                "pattern": "^[A-Za-z0-9-_]+$",
                "permission": "restricted_fields",
            },
            "consortia": {
                "title": "Consortia",
                "type": "array",
                "items": {"type": "string"},
            },
            "title": {"title": "Title", "type": "string"},
            "body": {"title": "Body", "type": "string"},
            "file": {"title": "File", "type": "string"},
            "section_type": {
                "title": "Section Type",
                "type": "string",
                "enum": ["Page Section", "Announcement"],
            },
            "status": {"type": "string", "permission": "restricted_fields"},
            "uuid": {"type": "string"},
            "aliases": {"type": "array", "items": {"type": "string"}},
            "schema_version": {"type": "string"},
            "submission_centers": {"type": "array", "serverDefault": "user_submission_centers"},
            "content": {"type": "string", "calculatedProperty": True},
            "options": {
                "type": "object",
                "properties": {
                    "filetype": {
                        "title": "File Type",
                        "type": "string",
                        "enum": ["md", "jsx"],
                    },
                    "collapsible": {
                        "title": "Collapsible",
                        "type": "boolean",
                    },
                },
            },
        },
    }


def test_load_config_accepts_inline_mapping_and_json_file(tmp_path):
    config = {"number_of_rows": 2, "section_type": "Page Section"}
    assert load_config(config) == config
    assert load_config(json.dumps(config)) == config
    config_path = tmp_path / "config.json"
    config_path.write_text(json.dumps(config), encoding="utf-8")
    assert load_config(config_path) == config

    with pytest.raises(StaticSectionWorkbookError, match="valid JSON"):
        load_config("{not-json")


def test_local_static_section_schema_is_resolved():
    schema = get_static_section_schema()
    properties = get_static_section_property_schemas(schema)
    assert properties["body"]["type"] == "string"
    assert properties["file"]["type"] == "string"
    assert properties["section_type"]["enum"]
    assert properties["options.filetype"]["enum"]
    assert "$merge" not in properties["body"]


def test_default_rows_and_blank_entry_columns(static_section_schema):
    workbook = generate_workbook(schema=static_section_schema)
    worksheet = workbook[STATIC_SECTION_SHEET_NAME]

    assert [cell.value for cell in worksheet[1]] == [
        "identifier",
        "consortia",
        "title",
        "body",
    ]
    assert worksheet.max_row == 2
    assert [cell.value for cell in worksheet[2]] == [None, None, None, None]


def test_configured_values_repeat_and_number_of_rows_is_not_a_column(static_section_schema):
    config = {
        "number_of_rows": 3,
        "section_type": "Page Section",
        "options.filetype": "jsx",
        "options.collapsible": False,
    }
    workbook = generate_workbook(config, schema=static_section_schema)
    worksheet = workbook.active

    assert "number_of_rows" not in [cell.value for cell in worksheet[1]]
    assert worksheet.max_row == 4
    headers = [cell.value for cell in worksheet[1]]
    values = [{header: worksheet.cell(row, column).value for column, header in enumerate(headers, 1)}
              for row in range(2, 5)]
    assert all(row["section_type"] == "Page Section" for row in values)
    assert all(row["options.filetype"] == "jsx" for row in values)
    assert all(row["options.collapsible"] is False for row in values)


def test_array_values_use_submitr_pipe_delimiter(static_section_schema):
    workbook = generate_workbook(
        {"consortia": ["consortium-a", "consortium-b"]},
        schema=static_section_schema,
    )
    assert workbook.active["B2"].value == "consortium-a | consortium-b"


def test_file_mode_is_explicit_and_exclusive(static_section_schema):
    workbook = generate_workbook(
        {"file": "/docs/public/section.md"},
        schema=static_section_schema,
    )
    assert [cell.value for cell in workbook.active[1]] == [
        "identifier",
        "consortia",
        "title",
        "file",
    ]
    assert workbook.active["D2"].value == "/docs/public/section.md"

    with pytest.raises(StaticSectionWorkbookError, match="mutually exclusive"):
        generate_workbook(
            {"body": "body", "file": "/section.md"},
            schema=static_section_schema,
        )


def test_config_schema_validation_and_read_only_fields(static_section_schema):
    with pytest.raises(StaticSectionWorkbookError, match="Unknown"):
        validate_config({"not_a_field": "value"}, static_section_schema)
    with pytest.raises(StaticSectionWorkbookError, match="read-only"):
        validate_config({"status": "open"}, static_section_schema)
    with pytest.raises(StaticSectionWorkbookError, match="read-only"):
        validate_config({"submission_centers": ["center"]}, static_section_schema)
    with pytest.raises(StaticSectionWorkbookError, match="not one of"):
        validate_config({"section_type": "Not a section"}, static_section_schema)
    with pytest.raises(StaticSectionWorkbookError, match="boolean"):
        validate_config({"options.collapsible": "false"}, static_section_schema)
    with pytest.raises(StaticSectionWorkbookError, match="positive integer"):
        validate_config({"number_of_rows": 0}, static_section_schema)


def test_identifier_is_writable_for_this_admin_workbook(static_section_schema):
    config = {"identifier": "section-one"}
    assert validate_config(config, static_section_schema) == config


def test_duplicate_fixed_identifier_is_rejected(static_section_schema):
    with pytest.raises(StaticSectionWorkbookError, match="duplicated"):
        validate_config(
            {"number_of_rows": 2, "identifier": "same-id"},
            static_section_schema,
        )


def test_workbook_structure_comments_and_output_naming(tmp_path, static_section_schema):
    output = write_static_section_workbook(
        tmp_path,
        {"section_type": "Page Section"},
        schema=static_section_schema,
    )
    assert output == Path(tmp_path, STATIC_SECTION_WORKBOOK_FILENAME)
    workbook = openpyxl.load_workbook(output)
    assert workbook.sheetnames == [STATIC_SECTION_SHEET_NAME]
    assert workbook.active["A1"].comment is not None
    assert "Type:" in workbook.active["A1"].comment.text
    assert get_output_path(tmp_path / "custom.xlsx") == tmp_path / "custom.xlsx"
    assert get_static_section_columns({}, static_section_schema) == [
        "identifier",
        "consortia",
        "title",
        "body",
    ]
