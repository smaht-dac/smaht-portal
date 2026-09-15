"""Generate a StaticSection entry workbook for administrative ingestion.

This command deliberately does not share the submission-schema lookup used by
``write_submission_spreadsheets``.  StaticSection is an admin-only item and
this template is generated from the portal's checked-in, locally resolved
schema so that it remains useful without portal credentials.
"""

from __future__ import annotations

import argparse
import json
from collections.abc import Mapping
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Union

import openpyxl
from jsonschema import Draft202012Validator
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from snovault import load_schema

STATIC_SECTION_SCHEMA = "encoded:schemas/static_section.json"
STATIC_SECTION_SHEET_NAME = "StaticSection"
STATIC_SECTION_WORKBOOK_FILENAME = "static_section_workbook.xlsx"
NUMBER_OF_ROWS = "number_of_rows"
DEFAULT_ENTRY_COLUMNS = ("identifier", "consortia", "title", "body")

# These fields either cannot be supplied by submitr or are populated by the
# server.  Keep the names here as a safeguard even when a schema version does
# not expose one of them in its properties.
READ_ONLY_FIELDS = frozenset(
    {
        "@id",
        "@type",
        "accession",
        "aliases",
        "content",
        "date_created",
        "filetype",
        "last_modified",
        "modified",
        "schema_version",
        "status",
        "submitted",
        "submitted_by",
        "submission_centers",
        "uuid",
    }
)

ConfigValue = Union[str, Path, Mapping[str, Any], None]


class StaticSectionWorkbookError(ValueError):
    """Raised when the workbook configuration is invalid."""


def load_config(config: ConfigValue) -> Dict[str, Any]:
    """Load a JSON object from an inline value, file, mapping, or ``None``.

    The mapping form makes the public builder easy to use from Python, while
    the command line accepts either a JSON object or a path to a JSON file.
    """
    if config is None:
        return {}
    if isinstance(config, Mapping):
        result = dict(config)
    else:
        if isinstance(config, Path):
            config_text = config.read_text(encoding="utf-8")
        elif isinstance(config, str):
            config_text = config.strip()
            config_path = Path(config_text)
            if config_text and not config_text.startswith("{") and config_path.is_file():
                config_text = config_path.read_text(encoding="utf-8")
        else:
            raise StaticSectionWorkbookError("Configuration must be a JSON object")
        try:
            result = json.loads(config_text or "{}")
        except json.JSONDecodeError as error:
            raise StaticSectionWorkbookError(
                f"Configuration is not valid JSON: {error.msg}"
            ) from error
    if not isinstance(result, dict):
        raise StaticSectionWorkbookError("Configuration must be a JSON object")
    return result


def get_static_section_schema() -> Dict[str, Any]:
    """Return the resolved local StaticSection schema."""
    return load_schema(STATIC_SECTION_SCHEMA)


def _flatten_schema_properties(
    properties: Mapping[str, Any], prefix: str = ""
) -> Iterable[tuple[str, Dict[str, Any]]]:
    """Yield scalar and array schema properties using dotted object paths."""
    for name, property_schema in properties.items():
        if name.startswith("$") or not isinstance(property_schema, Mapping):
            continue
        path = f"{prefix}.{name}" if prefix else name
        nested_properties = property_schema.get("properties")
        if nested_properties and isinstance(nested_properties, Mapping):
            yield from _flatten_schema_properties(nested_properties, path)
        else:
            yield path, dict(property_schema)


def get_static_section_property_schemas(
    schema: Optional[Mapping[str, Any]] = None,
) -> Dict[str, Dict[str, Any]]:
    """Return local schema properties addressable by workbook/config keys."""
    schema = schema or get_static_section_schema()
    properties = schema.get("properties")
    if not isinstance(properties, Mapping):
        raise StaticSectionWorkbookError(
            "The local StaticSection schema does not define properties"
        )
    property_schemas = dict(_flatten_schema_properties(properties))
    unresolved = [
        name for name, property_schema in property_schemas.items() if "$merge" in property_schema
    ]
    if unresolved:
        raise StaticSectionWorkbookError(
            "The local StaticSection schema contains unresolved properties: "
            + ", ".join(sorted(unresolved))
        )
    return property_schemas


def _is_read_only_property(name: str, property_schema: Mapping[str, Any]) -> bool:
    """Return whether a schema property is server-generated or calculated."""
    top_level_name = name.split(".", 1)[0]
    # The portal schema marks identifier as restricted because it is protected
    # on ordinary item endpoints, but it is the required user entry field for
    # this admin-only submitr workbook.
    if top_level_name == "identifier":
        return False
    if top_level_name in READ_ONLY_FIELDS:
        return True
    if "calculatedProperty" in property_schema or "serverDefault" in property_schema:
        return True
    if property_schema.get("permission") == "restricted_fields":
        return True
    excluded_from = property_schema.get("exclude_from", [])
    if isinstance(excluded_from, str):
        excluded_from = [excluded_from]
    return "FFedit-create" in excluded_from


def _validate_property_value(
    name: str, value: Any, property_schema: Mapping[str, Any]
) -> None:
    """Validate one configured value against its local property schema."""
    validator = Draft202012Validator(dict(property_schema))
    error = next(iter(validator.iter_errors(value)), None)
    if error is not None:
        raise StaticSectionWorkbookError(
            f"Invalid value for '{name}': {error.message}"
        )


def validate_config(
    config: Mapping[str, Any],
    schema: Optional[Mapping[str, Any]] = None,
) -> Dict[str, Any]:
    """Validate and return a StaticSection workbook configuration.

    Configured values are fixed template values and are checked individually
    against their corresponding schema property.  Entry columns are allowed
    to remain blank because the workbook is an entry form rather than a
    complete StaticSection item.
    """
    if not isinstance(config, Mapping):
        raise StaticSectionWorkbookError("Configuration must be a JSON object")
    config = dict(config)
    property_schemas = get_static_section_property_schemas(schema)

    number_of_rows = config.get(NUMBER_OF_ROWS, 1)
    if isinstance(number_of_rows, bool) or not isinstance(number_of_rows, int):
        raise StaticSectionWorkbookError("number_of_rows must be a positive integer")
    if number_of_rows < 1:
        raise StaticSectionWorkbookError("number_of_rows must be a positive integer")

    configured_names = set(config) - {NUMBER_OF_ROWS}
    unknown_names = configured_names - set(property_schemas)
    for name in sorted(unknown_names):
        if name.split(".", 1)[0] in READ_ONLY_FIELDS:
            raise StaticSectionWorkbookError(f"Field '{name}' is read-only")
        if name == "options":
            raise StaticSectionWorkbookError(
                "Field 'options' must use dotted keys such as 'options.filetype'"
            )
        raise StaticSectionWorkbookError(f"Unknown StaticSection field '{name}'")

    for name in sorted(configured_names):
        property_schema = property_schemas[name]
        if _is_read_only_property(name, property_schema):
            raise StaticSectionWorkbookError(f"Field '{name}' is read-only")
        _validate_property_value(name, config[name], property_schema)

    if "body" in config and "file" in config:
        raise StaticSectionWorkbookError(
            "StaticSection schema requires 'body' and 'file' to be mutually exclusive"
        )
    if "identifier" in config and number_of_rows > 1:
        raise StaticSectionWorkbookError(
            "A fixed identifier would be duplicated across multiple rows; "
            "omit identifier or use number_of_rows=1"
        )
    return config


def get_static_section_columns(
    config: Mapping[str, Any],
    schema: Optional[Mapping[str, Any]] = None,
) -> List[str]:
    """Return columns in entry-form order for a validated configuration."""
    config = validate_config(config, schema)
    columns = list(DEFAULT_ENTRY_COLUMNS)
    if "file" in config:
        columns.remove("body")
        columns.append("file")
    for name in config:
        if name != NUMBER_OF_ROWS and name not in columns:
            columns.append(name)
    return columns


def _serialize_cell_value(value: Any, property_schema: Mapping[str, Any]) -> Any:
    """Convert configured JSON values to values understood by submitr Excel."""
    if property_schema.get("type") != "array":
        return value
    # submitr uses a pipe-delimited cell for arrays.  The StaticSection schema
    # currently has string arrays, but str() keeps this safe for future scalar
    # array item schemas without changing workbook cell semantics.
    return " | ".join(str(item) for item in value)


def _schema_type(property_schema: Mapping[str, Any]) -> str:
    """Return a readable schema type for a header comment."""
    property_type = property_schema.get("type", "")
    if property_type == "array":
        item_type = property_schema.get("items", {}).get("type", "value")
        return f"array of {item_type}s"
    if isinstance(property_type, list):
        return " or ".join(str(item) for item in property_type)
    return str(property_type)


def _get_comment_text(
    name: str, property_schema: Mapping[str, Any], schema: Mapping[str, Any]
) -> str:
    """Build useful type and constraint information for a header comment."""
    lines = []
    if property_schema.get("title"):
        lines.append(property_schema["title"])
    if property_schema.get("description"):
        lines.append(f"Description: {property_schema['description']}")
    value_type = _schema_type(property_schema)
    if value_type:
        lines.append(f"Type: {value_type}")
    if property_schema.get("type") == "array":
        lines.append("Multiple values: separate values with |")
    if property_schema.get("enum"):
        lines.append("Options: " + " | ".join(str(item) for item in property_schema["enum"]))
    if property_schema.get("pattern"):
        lines.append(f"Pattern: {property_schema['pattern']}")
    if name in schema.get("required", []):
        lines.append("Required: Yes")
    else:
        lines.append("Required: No")
    if name == "body":
        lines.append("Do not fill file in the same row")
    elif name == "file":
        lines.append("Do not fill body in the same row")
    return "\n".join(lines)


def _write_header(
    worksheet: openpyxl.worksheet.worksheet.Worksheet,
    column: int,
    name: str,
    property_schema: Mapping[str, Any],
    schema: Mapping[str, Any],
) -> None:
    """Write and format one workbook header."""
    cell = worksheet.cell(row=1, column=column, value=name)
    cell.font = Font(name="Arial", size=10, bold=True)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.comment = Comment(
        _get_comment_text(name, property_schema, schema),
        "smaht-portal",
    )
    width = max(13, min(60, len(name) + 2))
    worksheet.column_dimensions[get_column_letter(column)].width = width


def generate_workbook(
    config: ConfigValue = None,
    schema: Optional[Mapping[str, Any]] = None,
) -> openpyxl.Workbook:
    """Generate a one-sheet StaticSection entry workbook."""
    schema = schema or get_static_section_schema()
    config = validate_config(load_config(config), schema)
    property_schemas = get_static_section_property_schemas(schema)
    columns = get_static_section_columns(config, schema)

    workbook = openpyxl.Workbook()
    worksheet = workbook.active
    worksheet.title = STATIC_SECTION_SHEET_NAME
    worksheet.freeze_panes = "A2"

    for column, name in enumerate(columns, start=1):
        _write_header(worksheet, column, name, property_schemas[name], schema)

    number_of_rows = config.get(NUMBER_OF_ROWS, 1)
    for row in range(2, number_of_rows + 2):
        for column, name in enumerate(columns, start=1):
            value = config.get(name)
            if name in config:
                value = _serialize_cell_value(value, property_schemas[name])
            cell = worksheet.cell(row=row, column=column, value=value)
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=name in {"body", "file"})

    return workbook


def get_output_path(output: Union[str, Path]) -> Path:
    """Resolve a directory or explicit ``.xlsx`` output path."""
    output = Path(output)
    if output.suffix.lower() == ".xlsx":
        return output
    return output / STATIC_SECTION_WORKBOOK_FILENAME


def write_static_section_workbook(
    output: Union[str, Path],
    config: ConfigValue = None,
    schema: Optional[Mapping[str, Any]] = None,
) -> Path:
    """Generate and save a StaticSection workbook, returning its path."""
    output_path = get_output_path(output)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = generate_workbook(config, schema=schema)
    workbook.save(output_path)
    return output_path


def main(argv: Optional[List[str]] = None) -> None:
    """Command-line entry point for the admin-only generator."""
    parser = argparse.ArgumentParser(
        description="Generate an admin-only StaticSection entry-form workbook"
    )
    parser.add_argument(
        "--output",
        required=True,
        help=(
            "Output .xlsx path or directory. Directories use "
            f"{STATIC_SECTION_WORKBOOK_FILENAME}."
        ),
    )
    parser.add_argument(
        "--config",
        help="Optional JSON object or path to a JSON configuration file",
    )
    args = parser.parse_args(argv)
    try:
        path = write_static_section_workbook(args.output, args.config)
    except (OSError, StaticSectionWorkbookError) as error:
        parser.error(str(error))
    print(f"StaticSection workbook written to: {path}")


if __name__ == "__main__":
    main()
